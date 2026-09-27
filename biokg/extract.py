"""Stage 1-3: ingest documents, repair encoding, segment checklist tables into records.

Output record (one per species row):
  source, waterbody, group, idx, raw_name, vn_name, hierarchy{rank: name},
  sites{site: '+'/'-'}, status{scheme: code}, campaigns, flags, raw_line_no
"""
import re
import subprocess
import unicodedata
from dataclasses import dataclass, field

import docx
from docx.table import Table

from .sources import DATA

# ---------------------------------------------------------------- repair
# Glyph substitutions observed in VAP proceedings PDFs (font re-encoding of Vietnamese vowels).
GLYPH_MAP = {"{": "à", "|": "á", "}": "â", "~": "ã", "[": "À", "\\": "Á"}
_GLYPH_RE = re.compile(r"(?<=[A-Za-zÀ-ỹĐđ])[{|}~\[\\]|[{|}~\[\\](?=[A-Za-zÀ-ỹĐđ])")


def repair_glyphs(text: str) -> tuple[str, int]:
    """Replace substituted glyphs that sit inside words. Returns (text, n_replacements)."""
    n = 0

    def sub(m):
        nonlocal n
        n += 1
        return GLYPH_MAP[m.group(0)]

    # citation brackets like "[12]" are not inside words, so they are left intact
    return _GLYPH_RE.sub(sub, text), n


def pdf_text(path) -> str:
    out = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, check=True)
    return unicodedata.normalize("NFC", out.stdout.decode("utf-8"))


def pdf_text_bbox(path, pt_per_char: float = 2.4) -> str:
    """Re-render a PDF from word bounding boxes on a fixed character grid.

    pdftotext -layout compresses whitespace line by line, so sparse tables (blank = absent)
    lose column alignment. Placing every word at round(xMin / pt_per_char) keeps all rows
    on the same horizontal scale as the header row.
    """
    import html
    out = subprocess.run(["pdftotext", "-bbox", str(path), "-"], capture_output=True, check=True)
    doc = out.stdout.decode("utf-8")
    pages = []
    for page in re.findall(r"<page[^>]*>(.*?)</page>", doc, re.S):
        words = [(float(a), float(b), float(c), html.unescape(w)) for a, b, c, w in
                 re.findall(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)"[^>]*>(.*?)</word>', page)]
        rows: list[list] = []
        for x, y, x2, w in sorted(words, key=lambda t: (t[1], t[0])):
            if rows and abs(rows[-1][0] - y) < 2.5:
                rows[-1][1].append((x, x2, w))
            else:
                rows.append([y, [(x, x2, w)]])
        lines = []
        for _, ws in rows:
            buf, prev_x2 = "", None
            for x, x2, w in sorted(ws):
                if prev_x2 is not None and x - prev_x2 < 4.5:   # same phrase: single space
                    buf += " " + w
                else:                                           # new column: snap to the grid
                    col = int(round(x / pt_per_char))
                    buf += " " * max(col - len(buf), 2 if buf else col) + w
                prev_x2 = x2
            lines.append(buf)
        pages.append("\n".join(lines))
    return unicodedata.normalize("NFC", "\n\f".join(pages))


# ---------------------------------------------------------------- ranks
RANK_WORDS = [  # (prefix as it appears, including glyph-loss variants, rank)
    ("Phân lớp", "subclass"), ("Phân ớ", "subclass"), ("Phân thứ lớp", "infraclass"),
    ("Phân họ", "subfamily"), ("Phân bộ", "suborder"),
    ("Giới", "kingdom"), ("Ngành", "phylum"), ("NGÀNH", "phylum"), ("Lớp", "class"), ("Lớ", "class"),
    ("LỚP", "class"), ("Bộ", "order"), ("BỘ", "order"), ("Order", "order"), ("Họ", "family"),
]
SUFFIX_RANK = [(r"idae$", "family"), (r"aceae$", "family"), (r"iformes$", "order"), (r"ales$", "order"),
               (r"(ida|optera|poda|cera)$", "order"), (r"(phyta|zoa|ophyta)$", "phylum"),
               (r"(phyceae|opsida|ceae)$", "class")]
LATIN_WORD = re.compile(r"\b[A-Z][A-Za-z]{3,}\b(?:\s+[a-z]{2,}\b)*")
ASCII_ONLY = re.compile(r"^[A-Za-z\- ]+$")


TAXON_SUFFIX = re.compile(r"(idae|aceae|iformes|ales|ida|phyta|phyceae|optera|poda|zoa|cera|ata|ia|ea)$", re.I)
PHYLA = {"Chordata", "Mollusca", "Arthropoda", "Annelida", "Rotifera", "Oomycota", "Magnoliophyta",
         "Cyanobacteria", "Chlorophyta", "Charophyta", "Ochrophyta", "Euglenozoa", "Miozoa", "Amoebozoa"}


def rank_of(text: str, marker: str = ""):
    """Return (rank, latin_name) for a hierarchy line, or None.

    The Latin part may be broken by glyph loss ('Arce   inida'), so a run of consecutive
    ASCII tokens is kept together and repaired later against a lexicon.
    """
    t = re.sub(r"\s+", " ", text.strip())
    rank = None
    for w, r in RANK_WORDS:
        if t.startswith(w + " "):
            rank, t = r, t[len(w) + 1:]
            break
    latin = None
    cands = [c.strip() for c in re.split(r"\s*[-–—]\s+|\s+[-–—]\s*|\(|\)", t) if c.strip()]
    for whole_only in (True, False):   # a fully Latin segment beats a Latin prefix of a mixed one
        for cand in cands:
            toks = cand.split()
            run = []
            for tok in toks:
                if not ASCII_ONLY.match(tok):
                    break
                run.append(tok)
            if not run or not run[0][0].isupper() or len("".join(run)) < 4:
                continue
            partial_ok = run[0].isupper() or TAXON_SUFFIX.search(run[-1]) or run[0] in PHYLA or bool(marker)
            if (len(run) == len(toks)) if whole_only else partial_ok:
                latin = " ".join(run)
                break
        if latin:
            break
    if latin is None:
        return None
    if rank is None and latin in PHYLA:
        rank = "phylum"
    if rank is None:
        for pat, r in SUFFIX_RANK:
            if re.search(pat, latin.split()[-1], re.I):
                rank = r
                break
    if rank is None and marker:          # numbering scheme of Vietnamese checklists: I, II = order; (1) = family
        rank = "family" if marker.startswith("(") else "order" if re.fullmatch(r"[IVXLC]+", marker) else None
    if rank is None:
        return None
    return rank, latin


# ---------------------------------------------------------------- records
@dataclass
class Record:
    source: str
    idx: int
    name_parts: list = field(default_factory=list)
    vn_name: str = ""
    hierarchy: dict = field(default_factory=dict)
    sites: dict = field(default_factory=dict)
    status: dict = field(default_factory=dict)
    campaigns: str = ""
    extra: list = field(default_factory=list)
    line_no: int = 0

    @property
    def raw_name(self):
        return re.sub(r"\s+", " ", " ".join(self.name_parts)).strip()


HEADER_NOISE = re.compile(
    r"BÁO CÁO KHOA HỌC|PHẦN \d|TẠP CHÍ|Tạp chí|pISSN|eISSN|^\s*(STT|Stt|TT)\b|Tên k|Điể|Điểm thu|"
    r"Tình trạng|Phân hạng|bảo tồn\s*$|^\s*\d{1,3}\s*$|^\s*\(\d\)\s+\(\d\)|IUCN\s*$|^\s*\(2025\)\s*$|"
    r"Nguồn gốc|xuất xứ|^\s*1\s+2\s+4\s+5|Tên Vi|Hoàng Đình Trung|^\s*(Loài|A\s+B\s+C)\s*$|DOI:|doi\.org|^\s*(ĐỘNG VẬT|THỰC VẬT|NẤM)\s*$")
TABLE_END = re.compile(r"^\s*(Bảng \d+\.|Ghi chú)|Tổng số|∑|\b(loài|taxon),\s*\d+\s*giống|^\s*\d+\s+loài\b")
STATUS_TOKEN = re.compile(r"^(LC|NT|VU|EN|CR|DD|NE|EW|EX|PL ?II|PL ?I|PL|II|NI|NII|NIB|PV|M|X)$")
SITE_TOKEN = re.compile(r"^[+\-xX]$")
VN_START = re.compile(r"^(Cá|Ốc|Nấm|Cây|Tôm|Cua|Trai|Hến|Rắn|Ếch|Cóc|Bèo|Rùa|Họ|Lươn|Chạch)\b")
VN_CHARS = re.compile(r"[ăâđêôơưàáảãạằắẳẵặầấẩẫậèéẻẽẹềếểễệìíỉĩịòóỏõọồốổỗộờớởỡợùúủũụừứửữựỳýỷỹỵ]", re.I)
VN_SPECIFIC = re.compile(r"[ăâđêôơưạảấầẩẫậắằẳẵặẹẻẽếềểễệỉịĩọỏốồổỗộớờởỡợụủũứừửữựỳỵỷỹ]|\bvà\b", re.I)
YEAR = re.compile(r"\b(1[7-9]\d\d|20[0-2]\d)\b")
MARKER = r"^\s*(?:[IVXLC]+|\(\d+\)|[A-Za-z]\d?)\s+"
RANK_ORDER = ["kingdom", "phylum", "class", "subclass", "infraclass", "order", "suborder", "family", "subfamily"]


def _fields(line):
    """Split a layout line into (start, end, text) fields separated by 2+ spaces."""
    return [(m.start(), m.end(), m.group(0)) for m in re.finditer(r"\S+(?: \S+)*", line)]


def _col_positions(line, pat):
    return [((m.start() + m.end()) / 2, m.group(0)) for m in re.finditer(pat, line)]


def _nearest(pos, cols, tol=4.5):
    best = min(cols, key=lambda c: abs(c[0] - pos))
    return best[1] if abs(best[0] - pos) <= tol else None


def _is_prose(line):
    words = line.split()
    return len(words) >= 9 and len(re.findall(r"\s{2,}", line.strip())) == 0 and \
        sum(bool(VN_CHARS.search(w)) for w in words) >= 4


TOTALS_ROW = re.compile(r"^\s*(?:\D{0,60}?)((?:\d{1,3}\s+){4,}\d{1,3})\s*$")


def segment_pdf(src: dict, text: str, meta: dict | None = None) -> list[Record]:
    lines = text.splitlines()
    start = next(i for i, l in enumerate(lines) if re.search(src["table"], l))
    site_cols, status_cols = [], []
    hier: dict = {}
    recs: list[Record] = []
    pending: list[str] = []  # name fragments seen before their index line (wrapped rows)
    status_map = src.get("status_cols", {})
    first_site_x = 10 ** 6
    prose_run = 0

    for ln in range(start + 1, len(lines)):
        line = lines[ln]
        if not line.strip():
            continue
        mt = TOTALS_ROW.match(line)
        if recs and mt and not re.match(r"^\s*\d{1,3}\s+[A-Z]", line) and "M1" not in line:
            if meta is not None:             # per-site (or per-campaign) richness printed by the authors
                meta["totals_row"] = [int(x) for x in mt.group(1).split()]
            break
        if TABLE_END.search(line):
            if meta is not None:             # totals may sit on this line or wrap onto the next two
                for cand in lines[ln:ln + 3]:
                    if m2 := re.search(r"((?:\d{1,3}\s+){4,}\d{1,3})\s*$", cand):
                        meta["totals_row"] = [int(x) for x in m2.group(1).split()]
                        break
            break
        sc = _col_positions(line, r"\bM\d{1,2}\b")
        if len(sc) >= 3:
            site_cols, first_site_x = sc, min(p for p, _ in sc) - 3
            continue
        stc = _col_positions(line, "|".join(r"(?<!\S)" + re.escape(k) + r"(?!\S)" for k in status_map)) \
            if status_map else []
        if len(stc) >= 3:
            status_cols = [(p, status_map.get(t, t)) for p, t in stc]
            continue
        lone_idx = pending and re.match(r"^\s*\d{1,3}\s*$", line)   # index row whose name wrapped around it
        if HEADER_NOISE.search(line) and not re.match(r"^\s*\d{1,3}\s+[A-Z][a-z]", line) and not lone_idx:
            continue
        prose_run = prose_run + 1 if _is_prose(line) else 0
        if prose_run >= 2 and recs:      # two consecutive prose lines: the table is over
            break
        if prose_run:                    # single prose line = running page title
            continue

        fl = _fields(line)
        m_idx = re.match(r"^\s*(\d{1,3})(?:\s+|$)(?!\d)", line)
        if m_idx:
            rec = Record(src["id"], int(m_idx.group(1)), hierarchy=dict(hier), line_no=ln + 1)
            b = m_idx.end()
            _fill(rec, [(s + b, e + b, t) for s, e, t in _fields(line[b:])], site_cols, status_cols, first_site_x)
            if pending and (not rec.name_parts or not re.match(r"[A-Z][a-z]+ ", rec.raw_name)):
                rec.name_parts = pending + rec.name_parts
            pending = []
            recs.append(rec)
            continue

        left = "  ".join(t for s, e, t in fl if s < first_site_x)
        mk = re.match(r"^\s*([IVXLC]+|\(\d+\))\s+", line)
        r = rank_of(re.sub(MARKER, "", left), mk.group(1) if mk else "")
        looks_species = re.match(r"^\s*[A-Z][a-z]+\s+([a-z]{3,}|sp\.)", re.sub(MARKER, "", left))
        if r and not looks_species:
            rank, name = r
            hier = {k: v for k, v in hier.items() if RANK_ORDER.index(k) < RANK_ORDER.index(rank)}
            hier[rank] = name
            continue

        # continuation fragment: belongs to the next index row if that row has no name, else to the previous one
        nxt = next((lines[j] for j in range(ln + 1, min(ln + 3, len(lines))) if lines[j].strip()), "")
        frag = " ".join(t for s, e, t in fl if s < first_site_x)
        m_n = re.match(r"^\s*\d{1,3}(?=\s|$)", nxt)
        if not src.get("no_merge") and m_n and re.match(r"^[A-Z][a-z]+ [a-z]", frag) \
                and not re.search(r"[A-Z][a-z]+ [a-z]{2,}", nxt[m_n.end():]):
            pending.append(frag)
        elif recs:
            _fill(recs[-1], fl, site_cols, status_cols, first_site_x, continuation=True)
    return recs


def _fill(rec, fields, site_cols, status_cols, first_site_x, continuation=False):
    """Distribute the fields of one layout line over name / vn_name / sites / status / extra."""
    site_marks = []
    for s, e, t in fields:
        toks = t.split()
        if site_cols and all(SITE_TOKEN.match(x) for x in toks):
            site_marks += [(s + m.start(), m.group(0)) for m in re.finditer(r"\S", t)]
            continue
        if all(STATUS_TOKEN.match(x) for x in toks):
            key = _nearest((s + e) / 2, status_cols, tol=8) if status_cols else None
            key = key or "codes"
            rec.status[key] = (rec.status.get(key, "") + " " + t).strip()
            continue
        if VN_START.match(t) or (rec.vn_name and VN_SPECIFIC.search(t) and not YEAR.search(t) and not continuation):
            if rec.vn_name and not VN_START.match(t):
                rec.extra.append(t)          # column after the Vietnamese name (e.g. origin)
            else:
                rec.vn_name = (rec.vn_name + " " + t).strip()
            continue
        if continuation and VN_SPECIFIC.search(t) and not YEAR.search(t) and not re.search(r"[A-Z][a-z]+ [a-z]", t):
            # wrapped Vietnamese name or wrapped origin cell
            if s > 40 and rec.extra:
                rec.extra[-1] += " " + t
            elif rec.vn_name:
                rec.vn_name += " " + t
            continue
        m = re.search(r"\s(\d(?:,\d)*)$", t)
        if m and re.search(r"\d{4}", t):  # trailing campaign list after the year: "Wallich, 1864 1,2,3"
            rec.campaigns = m.group(1)
            t = t[: m.start()]
        rec.name_parts.append(t)
    if site_marks and site_cols:
        if len(site_marks) == len(site_cols):          # complete row: assign by order (robust to drift)
            pairs = zip([c for _, c in sorted(site_cols)], [t for _, t in sorted(site_marks)])
        else:                                         # sparse row (blank = absent): assign by position
            pairs = [(_nearest(p, site_cols), t) for p, t in site_marks]
        for site, t in pairs:
            if site:
                rec.sites[site] = "+" if t in "+xX" else "-"


# ---------------------------------------------------------------- docx
def docx_rows(path, table_index):
    d = docx.Document(path)
    tb = d.tables[table_index]
    rows = []
    for r in tb.rows:
        cells, prev = [], None
        for c in r.cells:
            if c._tc is prev:
                continue
            prev = c._tc
            cells.append(c.text.strip().replace("\n", " "))
        rows.append(cells)
    return rows


def segment_docx(src: dict) -> list[Record]:
    rows = docx_rows(DATA / src["file"], src["table_index"])
    header = next((r for r in rows if any(re.fullmatch(r"M\d+", c) for c in r)), None)
    status_hdr = next((r for r in rows if sum(bool(re.fullmatch(r"\(\d\)", c)) for c in r) >= 3), None)
    hier, recs = {}, []
    order = ["kingdom", "phylum", "class", "subclass", "infraclass", "order", "suborder", "family", "subfamily"]
    for r in rows:
        if r is header or r is status_hdr or not r:
            continue
        idx = r[0].strip()
        if re.fullmatch(r"\d{1,3}", idx):
            rec = Record(src["id"], int(idx), [r[1]], hierarchy=dict(hier))
            if len(r) > 2 and not re.fullmatch(r"[+\-xX]?", r[2]):
                rec.vn_name = r[2]
            if header:
                for i, c in enumerate(header):
                    if re.fullmatch(r"M\d+", c) and i < len(r) and r[i] in "+-xX" and r[i]:
                        rec.sites[c] = "+" if r[i] in "+xX" else "-"
            if status_hdr:
                for i, c in enumerate(status_hdr):
                    if re.fullmatch(r"\(\d\)", c) and i < len(r) and r[i]:
                        rec.status[src.get("status_cols", {}).get(c, c)] = r[i]
            recs.append(rec)
        elif len(r) > 1 and r[1]:
            res = rank_of(r[1])
            if res:
                rank, name = res
                hier = {k: v for k, v in hier.items() if order.index(k) < order.index(rank)}
                hier[rank] = name
    return recs


# ---------------------------------------------------------------- driver
def extract_source(src: dict):
    """Return (records, meta). meta: glyph_fixes, totals_row (if the table prints one)."""
    path = DATA / src["file"]
    meta = {"glyph_fixes": 0}
    if src["kind"] == "pdf":
        raw = pdf_text_bbox(path) if src.get("render") == "bbox" else pdf_text(path)
        text, meta["glyph_fixes"] = repair_glyphs(raw)
        return segment_pdf(src, text, meta), meta
    return segment_docx(src), meta


# ---------------------------------------------------------------- sampling sites
DMS = re.compile(r"(\d{2,3})(?:°|0|\s)\s?(\d{1,2})\s?[’'′]\s?(\d{1,2}(?:[.,]\d+)?)")


def _dms(m):
    d, mi, se = m.groups()
    return int(d) + int(mi) / 60 + float(se.replace(",", ".")) / 3600


def sites_from_lines(lines) -> dict:
    """{site: (lat, lon)}. Latitude/longitude are told apart by value range, because several
    source tables print the 'Kinh độ' header above latitude values."""
    out = {}
    for line in lines:
        m = re.match(r"^\s*(M\d{1,2})\b(.*)$", line)
        if not m:
            continue
        vals = [_dms(x) for x in DMS.finditer(m.group(2))]
        lat = next((v for v in vals if 8 <= v <= 24), None)
        lon = next((v for v in vals if 102 <= v <= 110), None)
        if lat and lon:
            out[m.group(1)] = (round(lat, 6), round(lon, 6))
    return out


def extract_sites(src: dict) -> dict:
    path = DATA / src["file"]
    if src["kind"] == "pdf":
        lines = repair_glyphs(pdf_text(path))[0].splitlines()
        st = next((i for i, l in enumerate(lines) if re.search(r"Bảng 1\.\s*V[ịi] trí", l)), None)
        return sites_from_lines(lines[st:st + 60]) if st is not None else {}
    d = docx.Document(path)
    for tb in d.tables[:2]:
        rows = []
        for r in tb.rows:
            rows.append("  ".join(dict.fromkeys(c.text.strip() for c in r.cells)))
        got = sites_from_lines(rows)
        if got:
            return got
    return {}
