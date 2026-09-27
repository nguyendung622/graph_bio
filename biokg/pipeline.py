"""End-to-end run: extract -> parse -> resolve -> reconcile -> tables for the knowledge graph.

Usage:  python -m biokg.pipeline
Outputs go to extracted/gialai/.
"""
import json
import math
import re
from collections import defaultdict
from pathlib import Path

import pandas as pd

from .extract import extract_sites, extract_source
from .names import parse
from .resolve import RANKS, GbifResolver, classify_match, hint_query, query_for, repair_with_lexicon
from .sources import APPENDIX, DATA, SOURCES

OUT = Path(__file__).resolve().parent.parent / "extracted" / "gialai"


def haversine_m(a, b):
    la1, lo1, la2, lo2 = map(math.radians, (*a, *b))
    h = math.sin((la2 - la1) / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin((lo2 - lo1) / 2) ** 2
    return 2 * 6371000 * math.asin(math.sqrt(h))


# ------------------------------------------------------------------ appendix matrices
def appendix_abundance():
    """Species x site x campaign counts for Ayun Ha phyto/zooplankton (8 campaigns, 2021-22).

    Captions sit either in the paragraph before a table or in the table's first row
    (the appendix merges some consecutive tables), so both are tracked.
    """
    import docx as _docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = _docx.Document(DATA / APPENDIX["file"])
    rows_out, cap = [], ""
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            t = Paragraph(el, d).text.strip()
            if t.startswith("Bảng"):
                cap = t
            continue
        if tag != "tbl":
            continue
        sites = None
        for r in _rows(Table(el, d)):
            if r and r[0].startswith("Bảng"):
                cap, sites = r[0], None
                continue
            if "ực vật nổi" not in cap and "ộng vật nổi" not in cap:
                continue
            if any(re.fullmatch(r"M\d+", c) for c in r):
                sites = [(i, c) for i, c in enumerate(r) if re.fullmatch(r"M\d+", c)]
                continue
            m = re.search(r"tháng\s*(\d{1,2})\s*năm\s*(\d{4})", cap)
            if not (sites and m and re.fullmatch(r"\d+", r[0].strip())):
                continue
            date = f"{m.group(2)}-{int(m.group(1)):02d}"
            group = "phytoplankton" if "ực vật nổi" in cap else "zooplankton"
            trait = r[-1] if len(r) > sites[-1][0] + 1 and not re.fullmatch(r"[\d.,*]*", r[-1]) else ""
            for i, s in sites:
                v = r[i].strip().replace(".", "").replace(",", ".") if i < len(r) else ""
                if v in ("", "0"):
                    continue
                val = -1.0 if v in ("*", "+") else float(v) if re.fullmatch(r"[\d.]+", v) else None
                if val is not None:
                    rows_out.append(dict(source="AH_PL", waterbody="Hồ Ayun Hạ", group=group, campaign=date,
                                         site=s, raw_name=r[1], value=val, trait=trait))
    return pd.DataFrame(rows_out)


def appendix_water_quality():
    """17 water-quality parameters x 11 sites x 8 campaigns (Bảng 1PL1-8PL1 of the appendix)."""
    import docx as _docx
    d = _docx.Document(DATA / APPENDIX["file"])
    out = []
    for tb in d.tables[:8]:
        rows = _rows(tb)
        cap = tb._element.getprevious()
        cap_txt = "".join(cap.itertext()) if cap is not None else ""
        m = re.search(r"tháng\s*(\d{1,2})\s*năm\s*(\d{4})", cap_txt)
        hdr = next(r for r in rows if "M1" in r)
        for r in rows[rows.index(hdr) + 1:]:
            for i, site in enumerate(hdr):
                if re.fullmatch(r"M\d+", site) and i < len(r):
                    v = r[i].replace(",", ".")
                    if re.fullmatch(r"[\d.]+", v):
                        out.append(dict(source="AH_PL", waterbody="Hồ Ayun Hạ", campaign=f"{m.group(2)}-{int(m.group(1)):02d}",
                                        site=site, parameter=r[1], value=float(v)))
    return pd.DataFrame(out)


def _rows(tb):
    out = []
    for r in tb.rows:
        cells, prev = [], None
        for c in r.cells:
            if c._tc is prev:
                continue
            prev = c._tc
            cells.append(c.text.strip().replace("\n", " "))
        out.append(cells)
    return out


# ------------------------------------------------------------------ main
def run():
    OUT.mkdir(parents=True, exist_ok=True)
    rec_rows, presence_rows, site_rows, meta_rows = [], [], [], []
    for src in SOURCES:
        recs, meta = extract_source(src)
        meta_rows.append(dict(source=src["id"], glyph_fixes=meta["glyph_fixes"],
                              totals_row=json.dumps(meta.get("totals_row"))))
        for r in recs:
            p = parse(r.raw_name)
            rec_rows.append(dict(
                source=src["id"], waterbody=src["waterbody"], group=src["group"], idx=r.idx,
                raw_name=r.raw_name, vn_name=r.vn_name, canonical=p.canonical, genus=p.genus,
                authorship=p.authorship, year=p.year, is_morphospecies=p.is_morphospecies,
                is_higher_taxon=p.is_higher_taxon, symbols=p.symbols, name_notes=";".join(p.notes),
                src_phylum=r.hierarchy.get("phylum"), src_class=r.hierarchy.get("class"),
                src_order=r.hierarchy.get("order"), src_family=r.hierarchy.get("family"),
                status=json.dumps(r.status, ensure_ascii=False), extra="; ".join(r.extra),
                campaigns=r.campaigns, n_sites_marked=len(r.sites),
                n_sites_present=sum(v == "+" for v in r.sites.values())))
            sites_here = dict(r.sites)
            if src.get("blank_means_absent"):
                for k in range(1, src["reported"]["sites"] + 1):
                    sites_here.setdefault(f"M{k}", "-")
            for site, v in sites_here.items():
                presence_rows.append(dict(source=src["id"], idx=r.idx, site=site, present=v == "+"))
        for site, (lat, lon) in extract_sites(src).items():
            site_rows.append(dict(source=src["id"], waterbody=src["waterbody"], site=site, lat=lat, lon=lon))

    rec = pd.DataFrame(rec_rows)
    ab = appendix_abundance()
    ab_names = ab[["raw_name"]].drop_duplicates()
    ab_names["canonical"] = ab_names.raw_name.map(lambda s: parse(s).canonical)

    # ---- GBIF resolution (checklists + appendix names share one cache)
    resolver = GbifResolver(OUT / "gbif_cache.json")
    parsed_all = {n: parse(n) for n in set(rec.raw_name) | set(ab.raw_name)}
    queries = {n: query_for(p) for n, p in parsed_all.items()}
    res = resolver.match_many(list(queries.values()))

    def gb(n):
        q = queries[n]
        return res[f"{q[0]}|{q[1]}"]

    # second pass: unresolved names are re-queried with the source's family/order/kingdom as context
    first = {i: classify_match(gb(row.raw_name), parsed_all[row.raw_name]) for i, row in rec.iterrows()}
    retry = {i: hint_query(parsed_all[row.raw_name], row.src_family, row.src_order, row.group)
             for i, row in rec.iterrows() if first[i] in ("none", "higherrank")}
    res2 = resolver.match_many(list(retry.values()))

    for col in ["gbif_key", "gbif_name", "gbif_rank", "gbif_status", "gbif_accepted", "gbif_match",
                "gbif_confidence", "gbif_pass"] + [f"gbif_{k}" for k in RANKS]:
        rec[col] = None
    for i, row in rec.iterrows():
        g, p = gb(row.raw_name), parsed_all[row.raw_name]
        rec.at[i, "gbif_pass"] = 1
        if i in retry:
            g2 = res2["|".join(retry[i])]
            if classify_match(g2, p) not in ("none", "higherrank"):
                g = g2
                rec.at[i, "gbif_pass"] = 2
        rec.at[i, "gbif_key"] = g.get("usageKey")
        rec.at[i, "gbif_name"] = g.get("canonicalName")
        rec.at[i, "gbif_rank"] = g.get("rank")
        rec.at[i, "gbif_status"] = g.get("status")
        rec.at[i, "gbif_accepted"] = g.get("species") if g.get("status", "").endswith("SYNONYM") else None
        rec.at[i, "gbif_match"] = classify_match(g, p)
        rec.at[i, "gbif_confidence"] = g.get("confidence")
        for k in RANKS:
            rec.at[i, f"gbif_{k}"] = g.get(k)

    # ---- glyph-loss repair of source hierarchy names, with a per-source lexicon from GBIF
    repairs = []
    for sid, grp in rec.groupby("source"):
        lex = set()
        for k in ("phylum", "class", "order", "family"):
            lex |= {x for x in grp[f"gbif_{k}"].dropna()}
        for col in ("src_phylum", "src_class", "src_order", "src_family"):
            for dmg in grp[col].dropna().unique():
                if " " in dmg:
                    fix = repair_with_lexicon(dmg, lex)
                    repairs.append(dict(source=sid, rank=col[4:], damaged=dmg, repaired=fix))
                    if fix:
                        rec.loc[(rec.source == sid) & (rec[col] == dmg), col] = fix
    pd.DataFrame(repairs).to_csv(OUT / "hierarchy_repairs.csv", index=False)

    # ---- source-vs-GBIF family agreement
    def fam_norm(x):
        return (x or "").strip().lower()
    rec["family_agrees"] = [None if not (isinstance(a, str) and isinstance(b, str)) else fam_norm(a) == fam_norm(b)
                            for a, b in zip(rec.src_family, rec.gbif_family)]

    # guard: a fuzzy match whose GBIF family is unrelated to the printed family is rejected
    import difflib
    sim = [difflib.SequenceMatcher(None, str(a).lower(), str(b).lower()).ratio()
           for a, b in zip(rec.src_family, rec.gbif_family)]
    bad = (rec.gbif_match == "fuzzy") & (rec.family_agrees == False) & (pd.Series(sim) < 0.8)  # noqa: E712
    rec.loc[bad, "gbif_match"] = "fuzzy_rejected"
    rec["family_similarity"] = [round(x, 3) for x in sim]

    rec.to_csv(OUT / "checklist_records.csv", index=False)
    pd.DataFrame(presence_rows).to_csv(OUT / "site_presence.csv", index=False)
    pd.DataFrame(meta_rows).to_csv(OUT / "extraction_meta.csv", index=False)

    # ---- appendix abundance with resolution
    ab["canonical"] = ab.raw_name.map(lambda s: parsed_all[s].canonical)
    # reuse the (possibly second-pass) resolution obtained for the same canonical name in the checklists
    best = rec[~rec.gbif_match.isin(["none", "higherrank", "fuzzy_rejected"])].drop_duplicates("canonical").set_index("canonical")
    ab["gbif_key"] = [best.at[c, "gbif_key"] if c in best.index else gb(s).get("usageKey")
                      for s, c in zip(ab.raw_name, ab.canonical)]
    ab["gbif_match"] = [best.at[c, "gbif_match"] if c in best.index else classify_match(gb(s), parsed_all[s])
                        for s, c in zip(ab.raw_name, ab.canonical)]
    ab.to_csv(OUT / "ayunha_abundance.csv", index=False)
    wq = appendix_water_quality()
    wq.to_csv(OUT / "ayunha_water_quality.csv", index=False)

    # ---- sites: align across sources by distance (codes are source-local)
    sites = pd.DataFrame(site_rows)
    clusters, uid = [], 0
    for i, s in sites.iterrows():
        for c in clusters:
            if c["waterbody"] == s.waterbody and haversine_m((c["lat"], c["lon"]), (s.lat, s.lon)) < 150:
                c["members"].append(i)
                break
        else:
            uid += 1
            clusters.append(dict(waterbody=s.waterbody, lat=s.lat, lon=s.lon, members=[i], sid=f"S{uid:03d}"))
    for c in clusters:
        for i in c["members"]:
            sites.at[i, "site_uid"] = c["sid"]
    sites.to_csv(OUT / "sites.csv", index=False)
    return rec, ab, sites


if __name__ == "__main__":
    rec, ab, sites = run()
    print(rec.groupby("source").size())
    print(rec.gbif_match.value_counts())
    print(sites.groupby("waterbody").site_uid.nunique())
